"""Parsing, cleaning, classification and date-normalisation tests (no network, no models)."""
from __future__ import annotations

from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_BREAK

from agents.timeutil import parse_datetime
from core.config import get_settings
from core.models import Page, Table
from ingestion.classify import classify_document, pdf_date
from ingestion.cleaning import clean_pages, fix_broken_lines, fix_split_emails
from ingestion.html_parser import parse_html
from ingestion.pipeline import _parse, list_bid_files, parse_file

BIDS = Path(__file__).resolve().parents[1] / "data" / "bids"
needs_bids = pytest.mark.skipif(not BIDS.exists(), reason="sample bids not present")


# ---- cleaning ----------------------------------------------------------------------------

def test_hyphenated_ids_survive_line_breaks():
    assert fix_broken_lines("bid JA-\n207652 due") == "bid JA-207652 due"


def test_wrapped_sentence_is_rejoined():
    assert fix_broken_lines("the vendor shall\nprovide devices") == "the vendor shall provide devices"


def test_split_email_is_rejoined():
    assert fix_split_emails("thawkins@treasurer.state.md .us") == "thawkins@treasurer.state.md.us"


def test_running_header_removed_after_first_page():
    pages = [Page(page_number=i, text=f"Request For Proposal 168884 Page {i} of 40\nbody text {i}\nmore {i}")
             for i in range(1, 6)]
    clean_pages(pages)
    assert "Request For Proposal" in pages[0].text
    assert all("Request For Proposal" not in p.text for p in pages[1:])
    assert all(f"body text {p.page_number}" in p.text for p in pages)


def test_form_table_detection():
    form = Table(rows=[["Agency POC Name:", "", "", "Tamaira Hawkins", "", ""], ["", "", "Phone:", "", "", "410"]])
    grid = Table(rows=[["Item", "Qty"], ["Laptop", "30"]])
    assert form.is_form and not grid.is_form
    assert form.form_lines()[0] == "Agency POC Name: Tamaira Hawkins"


# ---- classification ----------------------------------------------------------------------

@pytest.mark.parametrize("name,head,expected", [
    ("Addendum 2 RFP JA-207652.pdf", "", ("addendum", 2)),
    ("misc.pdf", "ADDENDUM No. 3\nRFP 123", ("addendum", 3)),
    ("Contract_Affidavit.pdf", "", ("affidavit", None)),
    ("Dell_Laptop_Specs.pdf", "", ("specs", None)),
    ("PORFP_-_Dell_Laptop_Final.pdf", "", ("rfp", None)),
    ("notice.html", "", ("bid_page", None)),
    ("random.pdf", "hello", ("other", None)),
])
def test_classify_document(name, head, expected):
    assert classify_document(Path(name), head) == expected


def test_pdf_date():
    assert pdf_date({"ModDate": "D:20240613175856-05'00'"}) == "2024-06-13"
    assert pdf_date({}) is None


# ---- dates -------------------------------------------------------------------------------

def test_cst_in_summer_is_cdt():
    iso, note = parse_datetime("July 9, 2024 at 2:00 PM CST")
    assert iso == "2024-07-09T14:00:00-05:00" and "CDT" in note


def test_same_instant_across_documents():
    # Addendum 2 (2 PM "CST" = CDT) and the BidNet page (3 PM EDT) are the same moment
    from datetime import datetime
    a, _ = parse_datetime("July 9, 2024 at 2:00 PM CST")
    b, _ = parse_datetime("07/09/2024 03:00 PM EDT")
    assert datetime.fromisoformat(a) == datetime.fromisoformat(b)


def test_oracle_style_date_without_zone():
    assert parse_datetime("27-JUN-2024 14:00:00") == ("2024-06-27T14:00:00", "no time zone stated")


def test_no_date():
    assert parse_datetime("within 45 days of award") == (None, None)


# ---- docx --------------------------------------------------------------------------------

def test_docx_pages_tables_and_classification(tmp_path):
    path = tmp_path / "Addendum No. 3 - Test.docx"
    d = docx.Document()
    d.add_paragraph("ADDENDUM No. 3")
    d.add_paragraph("The due date is extended to August 1, 2024.")
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate([["Item", "Qty"], ["Laptop", "100"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.save(path)
    doc = _parse(path, "Test", "0" * 64, get_settings())
    assert (doc.doc_type, doc.addendum_number) == ("addendum", 3)
    assert len(doc.pages) == 2
    assert "August 1, 2024" in doc.pages[0].text
    assert doc.pages[1].tables[0].rows == [["Item", "Qty"], ["Laptop", "100"]]


def test_unsupported_and_corrupt_files_do_not_crash(tmp_path):
    (tmp_path / "notes.xlsx").write_bytes(b"x")
    bad = tmp_path / "broken.pdf"
    bad.write_bytes(b"%PDF-1.4 not really a pdf")
    assert list_bid_files(tmp_path) == [bad]
    settings = get_settings().model_copy(update={"parsed_cache_dir": tmp_path / "cache"})
    assert parse_file(bad, "Test", settings) is None


# ---- sample bids -------------------------------------------------------------------------

@needs_bids
def test_bidnet_html_fields():
    path = next((BIDS / "Bid2").glob("*.html"))
    pages, fields, _ = parse_html(path)
    assert fields["Solicitation Number"] == "BPM044557"
    assert fields["Closing Date"].startswith("06/10/2024")
    assert any("Thawkins@treasurer.state.md.us" in v for v in fields.values())
    assert pages[0].text.startswith("Reference Number")
