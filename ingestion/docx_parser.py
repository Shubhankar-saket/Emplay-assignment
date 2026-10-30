"""Word (.docx) parsing: paragraphs and tables in document order.

.docx has no fixed pages; page numbers follow Word's explicit page breaks and the page breaks Word
recorded at last save (w:lastRenderedPageBreak), so citations match the printed document where possible.
"""
from __future__ import annotations

from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

from core.models import Page, Table

_BREAK_TAGS = (qn("w:lastRenderedPageBreak"), qn("w:br"))


def _page_breaks(par: Paragraph) -> int:
    n = 0
    for el in par._p.iter(*_BREAK_TAGS):
        if el.tag == qn("w:lastRenderedPageBreak") or el.get(qn("w:type")) == "page":
            n += 1
    return n


def _table(tbl: DocxTable) -> Table | None:
    rows = []
    for row in tbl.rows:
        cells = [c.text.replace("\n", " ").strip() for c in row.cells]
        # merged cells repeat their text across the span; keep one copy
        rows.append([c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1]])
    rows = [r for r in rows if any(r)]
    return Table(rows=rows) if len(rows) >= 2 else None


def parse_docx(path: Path) -> tuple[list[Page], dict]:
    """Return Pages (split at page breaks) and a metadata dict compatible with PDF metadata keys."""
    d = docx.Document(str(path))
    pages = [Page(page_number=1)]
    lines: list[list[str]] = [[]]
    for block in d.iter_inner_content():
        if isinstance(block, Paragraph):
            for _ in range(_page_breaks(block)):
                if lines[-1] or pages[-1].tables:  # ignore a break at the very top of a page
                    pages.append(Page(page_number=len(pages) + 1))
                    lines.append([])
            text = block.text.strip()
            if text:
                lines[-1].append(text)
        elif isinstance(block, DocxTable) and (t := _table(block)):
            pages[-1].tables.append(t)
            lines[-1].extend(" | ".join(r) for r in t.rows)  # keep table text searchable in order
    for page, page_lines in zip(pages, lines):
        page.text = "\n".join(page_lines)
    props = d.core_properties
    meta = {"Title": props.title or "",
            "ModDate": props.modified.strftime("D:%Y%m%d") if props.modified else None,
            "CreationDate": props.created.strftime("D:%Y%m%d") if props.created else None}
    return pages, meta
